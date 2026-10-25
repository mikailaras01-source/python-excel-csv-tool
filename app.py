from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ------------------------------------------------------------
# Seitenkonfiguration und Konstanten
# ------------------------------------------------------------
st.set_page_config(
    page_title="Excel- und CSV-Aufbereitungstool",
    page_icon="📊",
    layout="wide",
)

PFLICHTSPALTEN = [
    "Bestellung",
    "Produkt",
    "Menge",
    "Einzelpreis (€)",
]
MENGENSPALTE = "Menge"
PREISSPALTE = "Einzelpreis (€)"
UMSATZSPALTE = "Umsatz (€)"


# ------------------------------------------------------------
# Hilfsfunktionen
# ------------------------------------------------------------
def datei_einlesen(datei) -> pd.DataFrame:
    """Liest eine XLSX- oder CSV-Datei ein und gibt klare Fehler aus."""
    dateiname = datei.name.lower()

    if dateiname.endswith(".xlsx"):
        return pd.read_excel(datei)
    if dateiname.endswith(".csv"):
        # UTF-8 ist Standard; bei häufigen CSV-Dateien mit deutschem
        # Windows-Encoding wird als Fallback cp1252 versucht.
        try:
            datei.seek(0)
            return pd.read_csv(datei)
        except UnicodeDecodeError:
            datei.seek(0)
            return pd.read_csv(datei, encoding="cp1252")

    raise ValueError("Nicht unterstütztes Dateiformat.")


def excel_export_erstellen(daten: pd.DataFrame, bericht: pd.DataFrame) -> bytes:
    """Erstellt eine formatierte Excel-Datei mit Daten und Qualitätsbericht."""
    ausgabe = io.BytesIO()

    with pd.ExcelWriter(ausgabe, engine="openpyxl") as writer:
        daten.to_excel(
            writer,
            index=False,
            sheet_name="Zusammengefuehrte_Daten",
        )
        bericht.to_excel(
            writer,
            index=False,
            sheet_name="Datenqualitaet",
        )

        for arbeitsblatt in writer.sheets.values():
            arbeitsblatt.freeze_panes = "A2"
            arbeitsblatt.auto_filter.ref = arbeitsblatt.dimensions
            arbeitsblatt.sheet_view.showGridLines = False

            for zelle in arbeitsblatt[1]:
                zelle.font = Font(bold=True, color="FFFFFF")
                zelle.fill = PatternFill(
                    fill_type="solid",
                    fgColor="164A3A",
                )
                zelle.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )
            arbeitsblatt.row_dimensions[1].height = 28

            for spaltenzellen in arbeitsblatt.columns:
                spaltennummer = spaltenzellen[0].column
                buchstabe = get_column_letter(spaltennummer)
                max_breite = max(
                    (
                        len(str(zelle.value))
                        for zelle in spaltenzellen
                        if zelle.value is not None
                    ),
                    default=10,
                )
                arbeitsblatt.column_dimensions[buchstabe].width = min(
                    max(max_breite + 3, 12),
                    40,
                )

        datenblatt = writer.sheets["Zusammengefuehrte_Daten"]
        header_map = {
            zelle.value: zelle.column
            for zelle in datenblatt[1]
        }
        for spaltenname in [PREISSPALTE, UMSATZSPALTE]:
            if spaltenname in header_map:
                spaltenbuchstabe = get_column_letter(header_map[spaltenname])
                for zeile in range(2, datenblatt.max_row + 1):
                    datenblatt[f"{spaltenbuchstabe}{zeile}"].number_format = '#,##0.00 €'

        # Tabellenformat, sofern Datenzeilen vorhanden sind.
        if datenblatt.max_row >= 2 and datenblatt.max_column >= 1:
            tabelle = Table(
                displayName="Verkaufsdaten",
                ref=datenblatt.dimensions,
            )
            tabelle.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            datenblatt.add_table(tabelle)

        berichtsblatt = writer.sheets["Datenqualitaet"]
        if berichtsblatt.max_row >= 2 and berichtsblatt.max_column >= 1:
            tabelle_bericht = Table(
                displayName="Datenqualitaetsbericht",
                ref=berichtsblatt.dimensions,
            )
            tabelle_bericht.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            berichtsblatt.add_table(tabelle_bericht)

    ausgabe.seek(0)
    return ausgabe.getvalue()


# ------------------------------------------------------------
# Benutzeroberfläche
# ------------------------------------------------------------
st.title("Excel- und CSV-Aufbereitungstool")
st.write(
    "Lade mehrere Excel- oder CSV-Dateien hoch, um sie zusammenzuführen, "
    "auf Datenqualität zu prüfen und als formatierte Excel-Datei zu exportieren."
)

with st.expander("Welche Dateien kann ich hochladen?", expanded=False):
    st.markdown(
        "- Unterstützte Formate: `.xlsx` und `.csv`\n"
        "- Die Dateien müssen dieselben Spaltennamen in derselben Reihenfolge enthalten.\n"
        "- Für die Umsatzberechnung werden `Menge` und `Einzelpreis (€)` benötigt.\n"
        "- Fehlende Werte und ungültige Mengen oder Preise werden vor dem Download geprüft."
    )

hochgeladene_dateien = st.file_uploader(
    "Excel- oder CSV-Dateien auswählen",
    type=["xlsx", "csv"],
    accept_multiple_files=True,
    help="Du kannst mehrere Dateien gleichzeitig auswählen.",
)

if not hochgeladene_dateien:
    st.info("Bitte lade mindestens eine Excel- oder CSV-Datei hoch.")
    st.stop()

alle_daten: list[pd.DataFrame] = []
datei_fehler: list[str] = []
datei_zusammenfassung: list[dict] = []

with st.expander("Dateivorschau", expanded=True):
    for datei in hochgeladene_dateien:
        try:
            daten = datei_einlesen(datei)

            if daten.empty:
                datei_fehler.append(
                    f"{datei.name}: Die Datei enthält keine Datenzeilen."
                )
                st.error(f"**{datei.name}:** Keine Datenzeilen gefunden.")
                continue

            if len(daten.columns) == 0:
                datei_fehler.append(
                    f"{datei.name}: Es wurden keine Spalten gefunden."
                )
                st.error(f"**{datei.name}:** Keine Spalten gefunden.")
                continue

            # Spaltennamen als Text vereinheitlichen und äußere Leerzeichen entfernen.
            daten.columns = [str(spalte).strip() for spalte in daten.columns]

            if len(set(daten.columns)) != len(daten.columns):
                datei_fehler.append(
                    f"{datei.name}: Doppelte Spaltennamen gefunden."
                )
                st.error(
                    f"**{datei.name}:** Die Datei enthält doppelte Spaltennamen. "
                    "Bitte benenne diese vor dem Upload eindeutig."
                )
                continue

            st.markdown(f"**{datei.name}**")
            st.caption(
                f"{len(daten):,} Datenzeilen · {len(daten.columns)} Spalten".replace(",", " ")
            )
            st.dataframe(daten.head(5), use_container_width=True, hide_index=True)

            alle_daten.append(daten)
            datei_zusammenfassung.append(
                {
                    "Datei": datei.name,
                    "Datenzeilen": len(daten),
                    "Spalten": len(daten.columns),
                    "Status": "Gelesen",
                }
            )

        except Exception as fehler:
            datei_fehler.append(
                f"{datei.name}: Datei konnte nicht gelesen werden ({fehler})."
            )
            st.error(
                f"**{datei.name}:** Die Datei konnte nicht gelesen werden. "
                "Prüfe, ob sie beschädigt oder korrekt formatiert ist."
            )

if datei_fehler:
    st.error("Die Verarbeitung wurde angehalten. Bitte korrigiere die oben genannten Dateien.")
    with st.expander("Technische Details"):
        for meldung in datei_fehler:
            st.write(f"- {meldung}")
    st.stop()

if not alle_daten:
    st.warning("Es konnten keine gültigen Dateien eingelesen werden.")
    st.stop()

referenz_spalten = list(alle_daten[0].columns)
abweichende_dateien = [
    datei.name
    for datei, daten in zip(hochgeladene_dateien, alle_daten)
    if list(daten.columns) != referenz_spalten
]

if abweichende_dateien:
    st.error("Die Dateien haben unterschiedliche Spalten.")
    st.write("Alle Dateien müssen dieselben Spaltennamen in derselben Reihenfolge enthalten.")
    st.write("**Spalten der ersten Datei:** " + ", ".join(referenz_spalten))
    for dateiname, daten in zip(
        [datei.name for datei in hochgeladene_dateien],
        alle_daten,
    ):
        if list(daten.columns) != referenz_spalten:
            st.write(f"**{dateiname}:** " + ", ".join(daten.columns))
    st.stop()

fehlende_pflichtspalten = [
    spalte for spalte in PFLICHTSPALTEN
    if spalte not in referenz_spalten
]
if fehlende_pflichtspalten:
    st.error("Die Dateien enthalten nicht alle benötigten Spalten.")
    st.write("Fehlende Spalten: " + ", ".join(f"`{s}`" for s in fehlende_pflichtspalten))
    st.write("Erforderlich sind: " + ", ".join(f"`{s}`" for s in PFLICHTSPALTEN))
    st.stop()

# Daten zusammenführen und Berechnungen durchführen.
zusammengefuehrte_daten = pd.concat(alle_daten, ignore_index=True)

menge_original = zusammengefuehrte_daten[MENGENSPALTE]
preis_original = zusammengefuehrte_daten[PREISSPALTE]
menge_numerisch = pd.to_numeric(menge_original, errors="coerce")
preis_numerisch = pd.to_numeric(preis_original, errors="coerce")

fehlende_mengen = menge_original.isna() | menge_original.astype(str).str.strip().eq("")
fehlende_preise = preis_original.isna() | preis_original.astype(str).str.strip().eq("")
fehlende_bestellungen = (
    zusammengefuehrte_daten["Bestellung"].isna()
    | zusammengefuehrte_daten["Bestellung"].astype(str).str.strip().eq("")
)
fehlende_produkte = (
    zusammengefuehrte_daten["Produkt"].isna()
    | zusammengefuehrte_daten["Produkt"].astype(str).str.strip().eq("")
)

menge_nicht_numerisch = ~fehlende_mengen & menge_numerisch.isna()
menge_ungueltig = menge_numerisch.notna() & (
    (menge_numerisch <= 0) | ((menge_numerisch % 1) != 0)
)
preis_nicht_numerisch = ~fehlende_preise & preis_numerisch.isna()
preis_ungueltig = preis_numerisch.notna() & (preis_numerisch < 0)

# Nullpreise sind zulässig (z. B. Gratisartikel); negative Preise nicht.
umsatz_gueltig = (
    ~fehlende_mengen
    & ~fehlende_preise
    & menge_numerisch.notna()
    & (menge_numerisch > 0)
    & ((menge_numerisch % 1) == 0)
    & preis_numerisch.notna()
    & (preis_numerisch >= 0)
)
zusammengefuehrte_daten[UMSATZSPALTE] = (
    menge_numerisch * preis_numerisch
).where(umsatz_gueltig)

doppelte_bestellungen = (
    zusammengefuehrte_daten["Bestellung"].notna()
    & zusammengefuehrte_daten["Bestellung"].astype(str).str.strip().ne("")
    & zusammengefuehrte_daten["Bestellung"].duplicated(keep=False)
)

fehler_mask = (
    fehlende_mengen
    | fehlende_preise
    | fehlende_bestellungen
    | fehlende_produkte
    | menge_nicht_numerisch
    | menge_ungueltig
    | preis_nicht_numerisch
    | preis_ungueltig
)

# Bericht mit übersichtlichen Prüfergebnissen.
pruefungen = [
    ("Dateien erfolgreich gelesen", len(hochgeladene_dateien), "OK"),
    ("Datensätze insgesamt", len(zusammengefuehrte_daten), "Info"),
    ("Fehlende Bestellnummern", int(fehlende_bestellungen.sum()), "Fehler" if fehlende_bestellungen.any() else "OK"),
    ("Fehlende Produktnamen", int(fehlende_produkte.sum()), "Fehler" if fehlende_produkte.any() else "OK"),
    ("Fehlende Mengen", int(fehlende_mengen.sum()), "Fehler" if fehlende_mengen.any() else "OK"),
    ("Nichtnumerische Mengen", int(menge_nicht_numerisch.sum()), "Fehler" if menge_nicht_numerisch.any() else "OK"),
    ("Ungültige Mengen (<= 0 oder nicht ganzzahlig)", int(menge_ungueltig.sum()), "Fehler" if menge_ungueltig.any() else "OK"),
    ("Fehlende Einzelpreise", int(fehlende_preise.sum()), "Fehler" if fehlende_preise.any() else "OK"),
    ("Nichtnumerische Einzelpreise", int(preis_nicht_numerisch.sum()), "Fehler" if preis_nicht_numerisch.any() else "OK"),
    ("Negative Einzelpreise", int(preis_ungueltig.sum()), "Fehler" if preis_ungueltig.any() else "OK"),
    ("Zeilen mit doppelter Bestellnummer", int(doppelte_bestellungen.sum()), "Hinweis" if doppelte_bestellungen.any() else "OK"),
]
datenqualitaet = pd.DataFrame(pruefungen, columns=["Prüfung", "Anzahl", "Status"])

# Kennzahlen
gesamtzeilen = len(zusammengefuehrte_daten)
anzahl_dateien = len(hochgeladene_dateien)
anzahl_fehler = int(fehler_mask.sum())
anzahl_doppelte = int(doppelte_bestellungen.sum())
umsatz_summe = zusammengefuehrte_daten[UMSATZSPALTE].sum(min_count=1)

st.subheader("Übersicht")
k1, k2, k3, k4 = st.columns(4)
k1.metric("Dateien", anzahl_dateien)
k2.metric("Datensätze", f"{gesamtzeilen:,}".replace(",", " "))
k3.metric("Zeilen mit Datenfehlern", anzahl_fehler)
k4.metric(
    "Umsatz gültiger Zeilen",
    "—" if pd.isna(umsatz_summe) else f"{umsatz_summe:,.2f} €".replace(",", "X").replace(".", ",").replace("X", "."),
)

st.subheader("Zusammengeführte Daten")
st.caption("Die Tabelle enthält alle eingelesenen Datensätze. Ungültige Umsätze bleiben leer, damit fehlerhafte Werte nicht versehentlich summiert werden.")
st.dataframe(
    zusammengefuehrte_daten,
    use_container_width=True,
    hide_index=True,
)

st.subheader("Datenqualität")
if anzahl_fehler == 0:
    st.success("Alle Pflichtfelder sowie Mengen und Preise haben die erwarteten Werte.")
else:
    st.error(f"Es wurden {anzahl_fehler} Zeilen mit mindestens einem Datenfehler gefunden.")

if anzahl_doppelte:
    st.warning(
        f"{anzahl_doppelte} Zeilen haben eine doppelte Bestellnummer. "
        "Das ist ein Hinweis zur Prüfung und blockiert den Export nicht automatisch."
    )

st.dataframe(datenqualitaet, use_container_width=True, hide_index=True)

if fehler_mask.any():
    st.markdown("**Zeilen mit Datenfehlern**")
    fehler_details = zusammengefuehrte_daten.loc[fehler_mask].copy()
    fehler_details.insert(0, "Zeile in Ergebnis", fehler_details.index + 2)
    st.dataframe(fehler_details, use_container_width=True, hide_index=True)

# Download nur bei gültigen Pflichtdaten; doppelte Bestellnummern sind lediglich ein Hinweis.
if anzahl_fehler > 0:
    st.warning(
        "Der Excel-Download ist gesperrt. Bitte korrigiere zuerst die fehlenden "
        "Pflichtwerte, Mengen- oder Preisfehler."
    )
else:
    export_bytes = excel_export_erstellen(zusammengefuehrte_daten, datenqualitaet)
    st.download_button(
        label="Formatierte Excel-Datei herunterladen",
        data=export_bytes,
        file_name="Zusammengefuehrte_Daten.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=False,
    )

st.caption(
    f"Erstellt mit Excel- und CSV-Aufbereitungstool · "
    f"{datetime.now().strftime('%d.%m.%Y')}"
)
