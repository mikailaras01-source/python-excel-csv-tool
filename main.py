from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ------------------------------------------------------------
# Einstellungen
# ------------------------------------------------------------
PROJEKT_ORDNER = Path(__file__).resolve().parent
EINGABE_ORDNER = PROJEKT_ORDNER / "beispieldaten"
AUSGABE_ORDNER = PROJEKT_ORDNER / "ergebnisse"
AUSGABE_DATEI = AUSGABE_ORDNER / "auswertung.xlsx"

DATEIEN = [
    EINGABE_ORDNER / "januar.xlsx",
    EINGABE_ORDNER / "februar.xlsx",
]

PFLICHTSPALTEN = [
    "Bestellung",
    "Produkt",
    "Menge",
    "Einzelpreis (€)",
]


# ------------------------------------------------------------
# Hilfsfunktionen
# ------------------------------------------------------------
def datei_einlesen(pfad: Path) -> pd.DataFrame:
    if not pfad.exists():
        raise FileNotFoundError(
            f"Die Datei wurde nicht gefunden: {pfad}\n"
            "Bitte prüfe, ob sie im Ordner 'beispieldaten' liegt."
        )

    try:
        daten = pd.read_excel(pfad)
    except Exception as fehler:
        raise ValueError(
            f"Die Datei '{pfad.name}' konnte nicht gelesen werden: {fehler}"
        ) from fehler

    if daten.empty:
        raise ValueError(
            f"Die Datei '{pfad.name}' enthält keine Datenzeilen."
        )

    daten.columns = [str(spalte).strip() for spalte in daten.columns]

    if len(set(daten.columns)) != len(daten.columns):
        raise ValueError(
            f"Die Datei '{pfad.name}' enthält doppelte Spaltennamen."
        )

    fehlende_spalten = [
        spalte for spalte in PFLICHTSPALTEN
        if spalte not in daten.columns
    ]
    if fehlende_spalten:
        raise ValueError(
            f"In '{pfad.name}' fehlen folgende Pflichtspalten: "
            f"{', '.join(fehlende_spalten)}"
        )

    return daten


def excel_formatieren(arbeitsblatt, waehrungsspalten: list[str] | None = None) -> None:
    waehrungsspalten = waehrungsspalten or []
    arbeitsblatt.freeze_panes = "A2"
    arbeitsblatt.auto_filter.ref = arbeitsblatt.dimensions
    arbeitsblatt.sheet_view.showGridLines = False
    arbeitsblatt.row_dimensions[1].height = 28

    for zelle in arbeitsblatt[1]:
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = PatternFill(fill_type="solid", fgColor="164A3A")
        zelle.alignment = Alignment(vertical="center", wrap_text=True)

    for spaltenzellen in arbeitsblatt.columns:
        buchstabe = get_column_letter(spaltenzellen[0].column)
        max_breite = max(
            (
                len(str(zelle.value))
                for zelle in spaltenzellen
                if zelle.value is not None
            ),
            default=10,
        )
        arbeitsblatt.column_dimensions[buchstabe].width = min(
            max(max_breite + 3, 12),
            40,
        )

    header_map = {
        zelle.value: zelle.column
        for zelle in arbeitsblatt[1]
    }
    for spaltenname in waehrungsspalten:
        if spaltenname in header_map:
            buchstabe = get_column_letter(header_map[spaltenname])
            for zeile in range(2, arbeitsblatt.max_row + 1):
                arbeitsblatt[f"{buchstabe}{zeile}"].number_format = '#,##0.00 €'

    if arbeitsblatt.max_row >= 2 and arbeitsblatt.max_column >= 1:
        tabellenname = f"Tabelle{arbeitsblatt.title.replace('_', '')}"
        # Excel-Tabellennamen dürfen keine Leerzeichen enthalten.
        tabellenname = "".join(
            zeichen for zeichen in tabellenname
            if zeichen.isalnum() or zeichen == "_"
        )
        tabelle = Table(displayName=tabellenname, ref=arbeitsblatt.dimensions)
        tabelle.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        arbeitsblatt.add_table(tabelle)


# ------------------------------------------------------------
# Verarbeitung
# ------------------------------------------------------------
def main() -> None:
    AUSGABE_ORDNER.mkdir(parents=True, exist_ok=True)

    eingelesene_daten = []
    for pfad in DATEIEN:
        daten = datei_einlesen(pfad)
        daten["Quelldatei"] = pfad.name
        eingelesene_daten.append(daten)
        print(f"OK: {pfad.name} – {len(daten)} Datenzeilen gelesen.")

    referenz_spalten = [
        spalte for spalte in eingelesene_daten[0].columns
        if spalte != "Quelldatei"
    ]

    for pfad, daten in zip(DATEIEN, eingelesene_daten):
        spalten = [
            spalte for spalte in daten.columns
            if spalte != "Quelldatei"
        ]
        if spalten != referenz_spalten:
            raise ValueError(
                f"Die Spalten in '{pfad.name}' stimmen nicht mit den "
                "Spalten der ersten Datei überein."
            )

    alle_daten = pd.concat(eingelesene_daten, ignore_index=True)

    # Leere Pflichtwerte prüfen.
    fehlende_mask = alle_daten[PFLICHTSPALTEN].isna()
    for spalte in PFLICHTSPALTEN:
        if alle_daten[spalte].dtype == "object":
            fehlende_mask[spalte] = (
                fehlende_mask[spalte]
                | alle_daten[spalte].astype(str).str.strip().eq("")
            )

    if fehlende_mask.any().any():
        details = alle_daten.loc[fehlende_mask.any(axis=1), PFLICHTSPALTEN]
        raise ValueError(
            "Die Verarbeitung wurde abgebrochen, weil Pflichtwerte fehlen.\n"
            f"{details.to_string(index=False)}"
        )

    menge = pd.to_numeric(alle_daten["Menge"], errors="coerce")
    preis = pd.to_numeric(alle_daten["Einzelpreis (€)"], errors="coerce")

    mengenfehler = (
        menge.isna()
        | (menge <= 0)
        | ((menge % 1) != 0)
    )
    if mengenfehler.any():
        details = alle_daten.loc[mengenfehler, PFLICHTSPALTEN]
        raise ValueError(
            "Ungültige Mengen gefunden. Mengen müssen positive ganze Zahlen sein.\n"
            f"{details.to_string(index=False)}"
        )

    preisfehler = preis.isna() | (preis < 0)
    if preisfehler.any():
        details = alle_daten.loc[preisfehler, PFLICHTSPALTEN]
        raise ValueError(
            "Ungültige Einzelpreise gefunden. Preise müssen Zahlen größer oder gleich 0 sein.\n"
            f"{details.to_string(index=False)}"
        )

    # Werte als numerische Datentypen speichern.
    alle_daten["Menge"] = menge.astype(int)
    alle_daten["Einzelpreis (€)"] = preis
    alle_daten["Umsatz (€)"] = alle_daten["Menge"] * alle_daten["Einzelpreis (€)"]

    # Doppelte Bestellnummern melden, aber nicht automatisch als Fehler behandeln.
    doppelte = alle_daten["Bestellung"].duplicated(keep=False)
    if doppelte.any():
        print(
            "HINWEIS: Es gibt doppelte Bestellnummern. Bitte prüfe, ob diese "
            "in deinen Daten zulässig sind."
        )
        print(alle_daten.loc[doppelte, PFLICHTSPALTEN].to_string(index=False))
    else:
        print("OK: Alle Bestellnummern sind eindeutig.")

    produkt_auswertung = (
        alle_daten.groupby("Produkt", dropna=False)
        .agg(
            Gesamtmenge=("Menge", "sum"),
            Gesamtumsatz=("Umsatz (€)", "sum"),
        )
        .reset_index()
        .sort_values("Produkt")
    )

    gesamtumsatz = alle_daten["Umsatz (€)"].sum()

    bericht = pd.DataFrame(
        [
            ("Eingelesene Dateien", len(DATEIEN)),
            ("Anzahl Datensätze", len(alle_daten)),
            ("Anzahl unterschiedlicher Produkte", alle_daten["Produkt"].nunique()),
            ("Doppelte Bestellnummern (Zeilen)", int(doppelte.sum())),
            ("Gesamtumsatz (€)", round(float(gesamtumsatz), 2)),
        ],
        columns=["Kennzahl", "Wert"],
    )

    with pd.ExcelWriter(AUSGABE_DATEI, engine="openpyxl") as writer:
        alle_daten.to_excel(
            writer,
            sheet_name="Einzelverkaeufe",
            index=False,
        )
        produkt_auswertung.to_excel(
            writer,
            sheet_name="Produktauswertung",
            index=False,
        )
        bericht.to_excel(
            writer,
            sheet_name="Qualitaetsbericht",
            index=False,
        )

        excel_formatieren(
            writer.sheets["Einzelverkaeufe"],
            waehrungsspalten=["Einzelpreis (€)", "Umsatz (€)"],
        )
        excel_formatieren(
            writer.sheets["Produktauswertung"],
            waehrungsspalten=["Gesamtumsatz"],
        )
        excel_formatieren(writer.sheets["Qualitaetsbericht"])

    print("\nVerarbeitung erfolgreich.")
    print(f"Datensätze: {len(alle_daten)}")
    print(f"Gesamtumsatz: {gesamtumsatz:.2f} €")
    print(f"Ergebnisdatei: {AUSGABE_DATEI}")


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError) as fehler:
        raise SystemExit(f"\nFEHLER: {fehler}") from fehler
